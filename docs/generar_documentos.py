"""Genera los documentos de entrega en entrega/: Fase 1 (análisis) y Fase 2 (documentación del software).

Ejecutar después de generar_diagramas.py:   python docs/generar_documentos.py
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent.parent
DIAG = RAIZ / "docs" / "diagramas"
CAPT = RAIZ / "docs" / "capturas"
ENTREGA = RAIZ / "entrega"
AZUL = RGBColor(0x1F, 0x4E, 0x79)
FECHA = "24 de octubre de 2026"
REPO = "https://github.com/Nicolas0pm/guia9-visualizacion-datos"


# ---------- utilidades de formato ----------

def documento():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.59), Cm(27.94)          # carta
    s.left_margin = s.right_margin = Cm(2.5)
    s.top_margin = s.bottom_margin = Cm(2.5)
    normal = d.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for nivel, tam in ((1, 15), (2, 12.5), (3, 11.5)):
        h = d.styles[f"Heading {nivel}"]
        h.font.name = "Arial"
        h.font.size = Pt(tam)
        h.font.color.rgb = AZUL
        h.font.bold = True
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        h.paragraph_format.space_before = Pt(14 if nivel == 1 else 10)
        h.paragraph_format.space_after = Pt(6)
    for zoom in d.settings.element.iter(qn("w:zoom")):      # la plantilla de python-docx omite este atributo
        zoom.set(qn("w:percent"), "100")
    pie = s.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _campo(pie, "PAGE")
    return d


def _campo(parrafo, codigo):
    run = parrafo.add_run()
    for tipo, texto in (("begin", None), (None, codigo), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        run._r.append(el)


def parrafo(d, texto, negrita=False, alinear=None, cursiva=False, tam=None):
    """Admite **negrita** dentro del texto."""
    p = d.add_paragraph()
    partes = texto.split("**")
    for i, parte in enumerate(partes):
        if parte:
            r = p.add_run(parte)
            r.bold = negrita or i % 2 == 1
            r.italic = cursiva
            if tam:
                r.font.size = Pt(tam)
    if alinear:
        p.alignment = alinear
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def vinetas(d, items, estilo="List Bullet"):
    for it in items:
        p = d.add_paragraph(style=estilo)
        for i, parte in enumerate(it.split("**")):
            if parte:
                p.add_run(parte).bold = i % 2 == 1


def tabla(d, encabezado, filas, anchos):
    t = d.add_table(rows=1, cols=len(encabezado))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(encabezado):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10)
        _sombra(c, "1F4E79")
    for j, fila in enumerate(filas):
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = ""
            for k, parte in enumerate(str(v).split("**")):
                if parte:
                    r = celdas[i].paragraphs[0].add_run(parte)
                    r.bold = k % 2 == 1
                    r.font.size = Pt(10)
            if j % 2:
                _sombra(celdas[i], "EEF3F9")
    for fila in t.rows:
        for i, w in enumerate(anchos):
            fila.cells[i].width = Cm(w)
    d.add_paragraph()
    return t


def _sombra(celda, color):
    tc = celda._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), color)
    tc.append(sh)


def figura(d, ruta, pie, ancho=16):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ruta), width=Cm(ancho))
    c = d.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(pie)
    r.italic = True
    r.font.size = Pt(9)


def portada(d, titulo, subtitulo):
    for _ in range(3):
        d.add_paragraph()
    for texto, tam, neg in (("GUÍA DE LABORATORIO N.º 9", 13, True),
                            ("Inteligencia de Negocios Avanzada y Toma de Decisiones", 12, False),
                            ("Eje temático: Business Analytics e Inteligencia Artificial", 12, False)):
        parrafo(d, texto, negrita=neg, alinear=WD_ALIGN_PARAGRAPH.CENTER, tam=tam)
    d.add_paragraph()
    d.add_paragraph()
    parrafo(d, titulo, negrita=True, alinear=WD_ALIGN_PARAGRAPH.CENTER, tam=20)
    parrafo(d, subtitulo, alinear=WD_ALIGN_PARAGRAPH.CENTER, tam=13)
    for _ in range(5):
        d.add_paragraph()
    parrafo(d, "Integrantes:", negrita=True, alinear=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(3):
        parrafo(d, "______________________________   Código: __________", alinear=WD_ALIGN_PARAGRAPH.CENTER)
    d.add_paragraph()
    parrafo(d, "Docente: ______________________________", alinear=WD_ALIGN_PARAGRAPH.CENTER)
    parrafo(d, f"Sesión 11 · Semana 11 · {FECHA}", alinear=WD_ALIGN_PARAGRAPH.CENTER)
    parrafo(d, f"Repositorio: {REPO}", alinear=WD_ALIGN_PARAGRAPH.CENTER, tam=10)
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------- FASE 1 ----------

def fase1():
    d = documento()
    portada(d, "FASE 1: Interpretación de datos", "Encuesta sobre actividades en el tiempo libre (n = 100)")

    d.add_heading("1. Situación inicial y datos", 1)
    parrafo(d, "Una institución educativa encuestó a 100 estudiantes sobre la actividad que realizan en su tiempo "
               "libre. Se trabajó únicamente con los datos que entrega la guía, sin modificarlos. Como la suma de "
               "las seis categorías es 100 y la población encuestada también es 100, cada estudiante aparece en una "
               "sola categoría y **el número de estudiantes coincide con el porcentaje**.")
    tabla(d, ["Actividad", "N.º de estudiantes", "% de la población", "% acumulado (de mayor a menor)"],
          [["Ver series o películas", 28, "28 %", "28 %"], ["Escuchar música", 22, "22 %", "50 %"],
           ["Practicar deporte", 18, "18 %", "68 %"], ["Jugar videojuegos", 15, "15 %", "83 %"],
           ["Leer", 10, "10 %", "93 %"], ["Otras actividades", 7, "7 %", "100 %"],
           ["**Total**", "**100**", "**100 %**", ""]], [6, 3.3, 3.2, 4])
    parrafo(d, "**Validación y trazabilidad.** Población encuestada = 100; suma de las categorías = 28 + 22 + 18 + "
               "15 + 10 + 7 = 100. No hay categorías repetidas ni valores negativos o decimales. Los indicadores de "
               "este documento se calcularon con estos mismos datos, tanto a mano como con el software de la FASE 2 "
               "(sus pruebas automáticas verifican que la moda sea «Ver series o películas», el mínimo «Otras "
               "actividades», el promedio 16,67 y el acumulado 28-50-68-83-93-100 %).")

    d.add_heading("2. Visualización de los datos", 1)
    figura(d, DIAG / "f1_barras.png", "Figura 1. Gráfico de barras con los datos de la guía (generado con el software de la FASE 2).")

    d.add_heading("3. Preguntas de interpretación", 1)
    preguntas = [
        ("¿Cuál es la actividad más frecuente?",
         "**Ver series o películas**, con 28 estudiantes (28 %). Es la barra más alta del gráfico."),
        ("¿Cuál es la actividad menos frecuente?",
         "**Otras actividades**, con 7 estudiantes (7 %). Entre las actividades con nombre propio, la menos "
         "frecuente es **Leer**, con 10 estudiantes (10 %)."),
        ("¿Cuántos estudiantes prefieren ver series o películas?", "**28 estudiantes**, el 28 % de los encuestados."),
        ("¿Cuántos estudiantes más prefieren escuchar música que leer?",
         "22 − 10 = **12 estudiantes más**. Escuchar música tiene 2,2 veces los estudiantes de Leer."),
        ("¿Qué actividades son realizadas por más de 15 estudiantes?",
         "**Ver series o películas (28), Escuchar música (22) y Practicar deporte (18)**. Jugar videojuegos tiene "
         "exactamente 15, por lo que no es «más de 15». Estas tres actividades reúnen 68 estudiantes (68 %)."),
        ("¿Qué porcentaje de los estudiantes practica deporte?",
         "18 / 100 × 100 = **18 %**. Casi uno de cada cinco estudiantes."),
        ("¿Qué diferencias importantes observas entre las actividades?",
         "La actividad más frecuente (28) **cuadruplica** a la menos frecuente (7), con una diferencia de 21 "
         "estudiantes. Las dos primeras (series y música) reúnen **la mitad** de los estudiantes (50 %). "
         "Ver series casi **triplica** a leer (28 frente a 10). Practicar deporte (18) supera a jugar "
         "videojuegos (15) por 3 estudiantes. Las diferencias entre categorías vecinas son pequeñas (6, 4, 3, 5 "
         "y 3 estudiantes): la distribución baja de forma escalonada, sin saltos bruscos."),
        ("¿Qué conclusión general puedes obtener a partir de los datos?",
         "Las preferencias se concentran en el **entretenimiento audiovisual y musical**: series y música suman "
         "50 %. Practicar deporte ocupa un lugar intermedio importante (18 %), y la lectura y otras actividades son "
         "minoritarias (17 % entre las dos). Ninguna actividad tiene la mayoría: la más elegida no llega a un "
         "tercio de los estudiantes."),
        ("¿Qué información no podemos conocer a partir de este gráfico?",
         "No sabemos **cuánto tiempo** dedica cada estudiante a cada actividad, ni con qué frecuencia. Tampoco "
         "sabemos si la pregunta permitía una sola respuesta (la suma de 100 lo sugiere), ni la edad, el grado o el "
         "sexo de los estudiantes. No sabemos qué incluye «Otras actividades», si la música se escucha en una "
         "pantalla, ni los motivos de cada preferencia. Tampoco conocemos cómo se eligieron los 100 estudiantes, ni "
         "si representan a toda la institución, ni si los resultados cambian con el tiempo."),
    ]
    for i, (q, r) in enumerate(preguntas, 1):
        parrafo(d, f"**3.{i} {q}**")
        parrafo(d, r)

    parrafo(d, "**3.10 Si tuvieras que presentar estos resultados a la institución, ¿qué otro tipo de gráfico "
               "utilizarías? Desarrollo de la propuesta y justificación.**")
    parrafo(d, "Se propone un **gráfico de barras horizontales ordenado de mayor a menor, con la cantidad y el "
               "porcentaje escritos al final de cada barra** (Figura 2). Las razones son las siguientes:")
    vinetas(d, [
        "**Ordenar** las barras permite leer el ranking sin buscar: la primera barra es la actividad más frecuente "
        "y la última la menos frecuente.",
        "La **orientación horizontal** deja leer completos los nombres largos («Ver series o películas»), que en "
        "el gráfico vertical se tienen que partir en dos líneas.",
        "Escribir el **valor y el porcentaje** evita que el público tenga que estimarlos a ojo con el eje, lo cual "
        "es importante en una presentación a directivos.",
        "Compara longitudes sobre una misma base, que el ojo juzga mejor que los ángulos de un gráfico circular. "
        "Por ejemplo, 15 % frente a 18 % se distingue bien en barras y mal en una torta.",
    ])
    figura(d, DIAG / "f1_propuesta_horizontal.png", "Figura 2. Propuesta: barras horizontales ordenadas con porcentaje.")
    parrafo(d, "Como complemento se propone un **diagrama de Pareto** (Figura 3). Muestra que tres actividades "
               "concentran el 68 % de las respuestas y cuatro el 83 %, lo que ayuda a la institución a decidir dónde "
               "enfocar sus programas de tiempo libre. El gráfico circular (Figura 4) también es válido, porque las "
               "partes suman el 100 % de la población. Sin embargo, con seis categorías de valores parecidos se "
               "comparan peor, así que se recomienda solo para mostrar la proporción general.")
    figura(d, DIAG / "f1_pareto.png", "Figura 3. Diagrama de Pareto: frecuencia y porcentaje acumulado.", 15)
    figura(d, DIAG / "f1_circular.png", "Figura 4. Gráfico circular (dona) de la proporción de cada actividad.", 13)

    d.add_heading("4. Reto de pensamiento crítico", 1)
    parrafo(d, "**Afirmación:** «El gráfico muestra que ver series o películas es la actividad más popular. ¿Podemos "
               "concluir que los estudiantes prefieren pasar todo su tiempo libre frente a una pantalla?»")
    parrafo(d, "**Posición: no estamos de acuerdo.** Los datos no permiten esa conclusión, por cuatro razones que se "
               "apoyan solo en la información disponible:")
    vinetas(d, [
        "**Es la más popular, pero no es la mayoría.** Ver series o películas tiene 28 estudiantes: el **72 %** "
        "eligió otra actividad.",
        "**Ni sumando las actividades de pantalla se llega a la mitad.** Las únicas categorías que claramente usan "
        "pantalla son ver series (28) y jugar videojuegos (15), que suman **43 estudiantes (43 %)**. Las que "
        "claramente no la usan, practicar deporte (18) y leer (10), suman **28 %**. Escuchar música (22 %) y "
        "otras actividades (7 %) no se pueden clasificar con estos datos.",
        "**La encuesta mide una preferencia, no el tiempo.** Registra qué actividad eligió cada estudiante, no "
        "cuántas horas le dedica. Por eso no se puede hablar de «todo su tiempo libre»: un estudiante que eligió "
        "series puede también hacer deporte o leer.",
        "**Se generaliza de 100 estudiantes a «los estudiantes».** No sabemos si la muestra representa a toda la "
        "institución.",
    ])
    parrafo(d, "Lo que sí se puede afirmar es que **ver series o películas es la preferencia individual más común "
               "(28 %)**, y que las actividades claramente asociadas a pantallas son una parte importante (43 %) pero "
               "no mayoritaria de las preferencias.")

    d.add_heading("5. Producto final: conclusión", 1)
    parrafo(d, "La encuesta aplicada a 100 estudiantes muestra que **el dato más alto corresponde a «Ver series o "
               "películas», elegida por 28 estudiantes (28 %)**. Es la actividad favorita del grupo, aunque no tiene "
               "la mayoría: casi tres de cada cuatro estudiantes prefieren otra cosa. Le sigue escuchar música con 22 "
               "estudiantes, de modo que entre las dos actividades reúnen exactamente a la mitad de los encuestados.")
    parrafo(d, "Al comparar categorías, **escuchar música (22 estudiantes) supera a leer (10 estudiantes) por 12 "
               "estudiantes, es decir, 2,2 veces**. Del mismo modo, practicar deporte (18) supera a jugar videojuegos "
               "(15) por 3 estudiantes, lo que muestra que la actividad física sigue siendo más elegida que los "
               "videojuegos.")
    parrafo(d, "**El patrón identificado es una concentración de las preferencias en pocas actividades**. Al "
               "ordenarlas de mayor a menor, la cantidad baja de forma escalonada (28, 22, 18, 15, 10 y 7), y las tres "
               "primeras acumulan el 68 % de las respuestas. También predominan las actividades de entretenimiento "
               "audiovisual y musical sobre la lectura, que es la actividad con nombre propio menos elegida (10 %).")
    parrafo(d, "**En conclusión, sustentada en los datos, los estudiantes prefieren principalmente actividades de "
               "entretenimiento (series y música, 50 %), pero sus preferencias son diversas**: ninguna actividad supera "
               "el 28 % y el deporte reúne casi a uno de cada cinco. Para la institución, la baja preferencia por la "
               "lectura (10 %) es la principal oportunidad de intervención, y el interés por el deporte (18 %) es una "
               "fortaleza que se puede aprovechar. Para conocer cuánto tiempo dedican a cada actividad haría falta "
               "una nueva encuesta que pregunte horas y permita más de una respuesta.")

    d.add_heading("6. Referencias", 1)
    parrafo(d, "Escobar Mercado, M., Fernández Macías, E. y Bernardi, F. (2021). Análisis de datos con Stata (3.ª ed.). "
               "CIS — Centro de Investigaciones Sociológicas.")
    parrafo(d, "Universidad — Guía de laboratorio N.º 9: Business Analytics e Inteligencia Artificial (2026).")
    parrafo(d, f"Repositorio del proyecto (gráficas, datos y software de la FASE 2): {REPO}")
    d.save(ENTREGA / "Fase1_Interpretacion_de_datos.docx")


# ---------- FASE 2 ----------

def fase2():
    d = documento()
    portada(d, "FASE 2: Documentación del software", "Visualizador de Encuestas: aplicación de escritorio portable "
                                                     "para ingresar datos y generar gráficos")

    d.add_heading("1. Introducción", 1)
    parrafo(d, "El **Visualizador de Encuestas** es una aplicación de escritorio que permite ingresar de forma "
               "estructurada los datos de una encuesta, validarlos, guardarlos y generar automáticamente sus gráficos e "
               "indicadores. Trae cargados los datos de la FASE 1 (100 estudiantes y seis actividades), pero acepta "
               "cualquier otra encuesta de categorías y frecuencias. Está hecha en **Python con interfaz gráfica "
               "Tkinter** y se entrega como un **único archivo .exe que no requiere instalación**, ni Python, ni "
               "internet. No es una aplicación web, ni PWA, ni usa HTML, CSS o JavaScript, como exige la guía.")
    tabla(d, ["Elemento", "Valor"], [
        ["Nombre", "Visualizador de Encuestas 1.0"],
        ["Lenguaje", "Python 3.12"],
        ["Interfaz gráfica", "Tkinter / ttk (incluida en Python)"],
        ["Gráficos", "matplotlib 3.10 integrado en la ventana (FigureCanvasTkAgg)"],
        ["Almacenamiento", "SQLite (archivo datos/encuesta.db junto al ejecutable)"],
        ["Empaquetado", "PyInstaller 6.16 en modo --onefile --windowed"],
        ["Sistema operativo", "Windows 10/11 de 64 bits (el código fuente también corre en Linux y macOS)"],
        ["Pruebas", "23 pruebas automáticas con unittest + prueba de la interfaz con capturas"],
        ["Código fuente", REPO],
    ], [5, 11])

    d.add_heading("2. Requisitos", 1)
    d.add_heading("2.1 Requisitos funcionales", 2)
    tabla(d, ["ID", "Requisito", "Dónde se cumple"], [
        ["RF1", "Agregar una actividad con su número de estudiantes", "Formulario «Ingreso de datos», botón Agregar o Enter"],
        ["RF2", "Validar cada dato antes de guardarlo y explicar el error", "modelo.py; mensaje en rojo junto al campo"],
        ["RF3", "Actualizar y eliminar datos existentes", "Seleccionar fila → Actualizar / Eliminar (o tecla Supr)"],
        ["RF4", "Guardar los datos de forma permanente", "almacen.py → SQLite datos/encuesta.db"],
        ["RF5", "Graficar automáticamente el dato ingresado", "graficos.py; el gráfico se redibuja tras cada cambio y "
                                                           "resalta en naranja el dato nuevo"],
        ["RF6", "Ofrecer varios tipos de gráfico", "Barras, barras horizontales %, circular, Pareto, comparación"],
        ["RF7", "Mostrar indicadores (KPI)", "Tarjetas: total/cobertura, más y menos frecuente, promedio"],
        ["RF8", "Configurar título y población de la encuesta", "Recuadro «Encuesta», botón Aplicar"],
        ["RF9", "Importar y exportar datos en CSV", "Menú Archivo"],
        ["RF10", "Exportar el gráfico como imagen PNG", "Botón Exportar PNG o menú Archivo"],
        ["RF11", "Conservar la trazabilidad respecto a los datos originales",
         "Pestaña Historial, gráfico de comparación y opción Restaurar"],
    ], [1.4, 7.3, 7.3])
    d.add_heading("2.2 Requisitos no funcionales", 2)
    tabla(d, ["ID", "Requisito", "Solución"], [
        ["RNF1", "Portable, sin instalación", "Un solo .exe (PyInstaller --onefile); la base de datos se crea junto a él"],
        ["RNF2", "Interfaz gráfica de escritorio, no web", "Tkinter; no usa navegador, servidor ni internet"],
        ["RNF3", "Integridad de los datos", "Validación en modelo.py + restricciones CHECK/UNIQUE en SQLite + transacciones"],
        ["RNF4", "Usabilidad", "Un solo formulario, Enter para guardar, mensajes en lenguaje claro, ayuda en el menú"],
        ["RNF5", "Código documentado y probado", "Docstrings en cada módulo y función; 23 pruebas automáticas"],
    ], [1.6, 5.4, 9])

    d.add_heading("3. Selección de herramientas", 1)
    parrafo(d, "**Python** se eligió porque tiene bibliotecas maduras para datos y gráficos, y porque el mismo código "
               "sirve para las pruebas, para los diagramas de esta documentación y para la aplicación. **Tkinter** viene "
               "incluido en Python, así que no agrega dependencias. **matplotlib** produce gráficos de calidad "
               "de publicación y se integra en una ventana Tkinter. **SQLite** es una base de datos en un solo archivo, "
               "sin servidor, lo que la hace ideal para un programa portable. **PyInstaller** empaqueta el intérprete, "
               "las bibliotecas y el código en un solo .exe, para que el equipo de destino no necesite tener nada "
               "instalado.")

    d.add_heading("4. Arquitectura", 1)
    parrafo(d, "El software se organiza en capas. La interfaz no valida ni escribe por su cuenta: delega en el "
               "almacén, y este aplica las reglas de modelo.py antes de cada escritura. Los gráficos se generan con "
               "funciones independientes de la interfaz, por lo que se pueden probar y exportar sin abrir ventanas.")
    figura(d, DIAG / "01_arquitectura.png", "Figura 1. Arquitectura del software por capas.")
    parrafo(d, "Desde la perspectiva de Inteligencia de Negocios, la aplicación sigue el flujo fuente → ETL → "
               "almacenamiento → indicadores → visualización → decisión:")
    figura(d, DIAG / "02_flujo_bi.png", "Figura 2. Flujo de la información en la solución.")

    d.add_heading("5. Diagramas de diseño", 1)
    d.add_heading("5.1 Casos de uso", 2)
    figura(d, DIAG / "04_casos_de_uso.png", "Figura 3. Diagrama de casos de uso.", 14)
    d.add_heading("5.2 Clases y módulos", 2)
    figura(d, DIAG / "05_clases.png", "Figura 4. Diagrama de clases y módulos.")
    d.add_heading("5.3 Modelo de datos", 2)
    parrafo(d, "La base tiene tres tablas. **encuesta** guarda una sola fila con el título y la población. "
               "**categoria** guarda cada actividad, su cantidad actual, la cantidad original de la guía y su origen. "
               "**historial** registra cada cambio con fecha, acción y detalle (valor anterior → valor nuevo).")
    figura(d, DIAG / "06_modelo_datos.png", "Figura 5. Modelo de datos en SQLite.")
    d.add_heading("5.4 Diagrama de flujo del ingreso de datos", 2)
    parrafo(d, "El mismo algoritmo se entrega en pseudocódigo de PSeInt (docs/diagramas/agregar_dato.psc), que "
               "puede abrirse en PSeInt para dibujar su diagrama de flujo.")
    figura(d, DIAG / "03_flujo_agregar.png", "Figura 6. Diagrama de flujo para agregar un dato.", 12)

    d.add_heading("6. Ingreso, validación y gestión de datos", 1)
    parrafo(d, "**¿Qué datos se ingresan y para qué?** Pares (actividad, número de estudiantes) de una encuesta, "
               "más el título y la población encuestada. Con ellos se calculan los porcentajes, los indicadores y los "
               "gráficos que sirven para interpretar las preferencias del grupo.")
    parrafo(d, "**¿Cómo se garantiza que sean completos, válidos y consistentes?** Se aplican tres niveles de control. "
               "El primero es un filtro de teclado: los campos numéricos solo aceptan dígitos. El segundo son las reglas "
               "de modelo.py, que usan el formulario, la importación de CSV y el almacén. El tercero son las "
               "restricciones de la base de datos (CHECK, UNIQUE, NOT NULL). Cada escritura se hace en una transacción.")
    tabla(d, ["Campo / regla", "Validación", "Mensaje al usuario (ejemplo)"], [
        ["Actividad vacía", "Obligatoria, se quitan espacios repetidos", "Escriba el nombre de la actividad."],
        ["Largo y caracteres", "2 a 40 caracteres, al menos una letra, sin símbolos raros",
         "El nombre debe tener al menos 2 caracteres."],
        ["Duplicados", "No se repite, sin importar tildes ni mayúsculas (LEÉR = Leer)",
         "Ya existe la actividad «Leer». Selecciónela en la tabla para cambiarla."],
        ["N.º de estudiantes", "Entero, sin decimales ni letras, entre 0 y la población",
         "El número de estudiantes no puede ser mayor que 100."],
        ["Coherencia", "La suma de todas las actividades ≤ población encuestada",
         "Con este dato la suma sería 105 estudiantes y la encuesta tiene 100. Sobran 5."],
        ["Población", "Entero de 1 a 1.000.000 y nunca menor que lo ya registrado",
         "La población (90) no puede ser menor que los estudiantes ya registrados (100)."],
        ["Cantidad de actividades", "Máximo 15, para que el gráfico se lea bien", "Se admiten como máximo 15 actividades…"],
        ["Importar CSV", "Todo o nada: si una fila falla no se importa ninguna; indica la línea",
         "No se importó nada. Corrija el archivo: Línea 3: Escriba el nombre…"],
    ], [3.6, 6.2, 6.2])

    d.add_heading("7. Visualización automatizada", 1)
    parrafo(d, "Después de cada cambio, la ventana vuelve a leer la base de datos y redibuja la tabla, los indicadores "
               "y el gráfico elegido. El dato que se acaba de ingresar o la fila seleccionada se pinta en naranja. "
               "**¿Por qué cada tipo de visualización?**")
    tabla(d, ["Gráfico", "Para qué sirve"], [
        ["Barras verticales", "Comparar cantidades entre categorías; es el gráfico de la guía."],
        ["Barras horizontales ordenadas (%)", "Ranking de mayor a menor con cantidad y porcentaje; nombres largos legibles."],
        ["Circular (dona)", "Proporción de cada actividad sobre la población; muestra en gris lo que falta registrar."],
        ["Pareto", "Qué pocas actividades concentran la mayoría de las respuestas (línea de 80 %)."],
        ["Comparación con datos originales", "Dato de la guía frente al dato actual: trazabilidad de los cambios."],
    ], [5.5, 10.5])
    parrafo(d, "**¿Cómo facilita la interpretación?** Las etiquetas de valor evitan estimar a ojo y el orden muestra el "
               "ranking. Los porcentajes se calculan sobre la población, las tarjetas KPI resumen las respuestas a las "
               "preguntas de la FASE 1 (más y menos frecuente, total, promedio) y la nota «n = … de … estudiantes» deja "
               "claro sobre qué base se calcula cada porcentaje.")

    d.add_heading("8. Interfaz de usuario", 1)
    figura(d, CAPT / "05_agregado_barras.png", "Figura 7. Ventana principal tras agregar «Dibujar o pintar» (resaltado en naranja).")
    vinetas(d, [
        "**Encuesta:** título y población encuestada (denominador de los porcentajes).",
        "**Tarjetas KPI:** estudiantes registrados y cobertura, actividad más y menos frecuente, promedio.",
        "**Ingreso de datos:** actividad, número de estudiantes, botones Agregar, Actualizar, Eliminar y Limpiar; "
        "debajo se indica cuántos estudiantes faltan por asignar.",
        "**Datos registrados:** tabla con cantidad, porcentaje y origen (original, ingresado o importado).",
        "**Gráfico:** selector de tipo, exportación a PNG y barra de zoom y desplazamiento de matplotlib.",
        "**Historial de cambios:** cada operación con fecha, acción y detalle.",
    ])
    figura(d, CAPT / "02_supera_poblacion.png", "Figura 8. Validación: el dato superaría la población y no se guarda.", 14)
    figura(d, CAPT / "07_editado_pareto.png", "Figura 9. Edición de «Leer» (10 → 14) y diagrama de Pareto.", 14)
    figura(d, CAPT / "10_historial.png", "Figura 10. Historial de cambios (trazabilidad).", 14)

    d.add_heading("9. Documentación del código", 1)
    parrafo(d, "Cada módulo empieza con un docstring que explica su responsabilidad, y cada función pública "
               "documenta qué recibe, qué devuelve y qué error lanza. Estructura del proyecto:")
    tabla(d, ["Archivo", "Responsabilidad"], [
        ["app/main.py", "Punto de entrada. Ubica la base de datos junto al .exe (o en Documentos si la carpeta es de solo "
                        "lectura) y abre la ventana."],
        ["app/modelo.py", "Reglas del negocio: datos originales, validaciones, coherencia, indicadores, orden y % acumulado."],
        ["app/almacen.py", "Clase Almacen: esquema SQLite, CRUD, historial, restaurar, importar y exportar CSV."],
        ["app/graficos.py", "Función dibujar() con los 5 tipos de gráfico y exportar_png()."],
        ["app/interfaz.py", "Clase Aplicacion (ventana Tkinter) y clase Tarjeta (KPI)."],
        ["tests/test_app.py", "23 pruebas automáticas (validación, indicadores, almacén, CSV y gráficos)."],
        ["docs/generar_diagramas.py", "Regenera los diagramas y las gráficas de esta documentación."],
        ["construir_windows.bat", "Construye el .exe en cualquier Windows con Python (ejecuta las pruebas antes)."],
        [".github/workflows/construir-exe.yml", "Construye el .exe automáticamente en GitHub Actions."],
    ], [5.5, 10.5])
    parrafo(d, "Funciones principales:")
    tabla(d, ["Función / método", "Descripción"], [
        ["validar_nombre(texto, existentes, ignorar)", "Limpia el nombre y lanza ErrorValidacion si está vacío, es "
                                                       "muy corto o largo, tiene símbolos o ya existe."],
        ["validar_cantidad(texto, poblacion)", "Convierte a entero entre 0 y la población."],
        ["validar_coherencia(suma, poblacion)", "Impide que la suma de actividades supere la población."],
        ["calcular_indicadores(categorias, poblacion)", "Devuelve total, cobertura, promedio, más y menos frecuentes "
                                                        "(con empates)."],
        ["Almacen.agregar / actualizar / eliminar", "Validan, escriben en una transacción y registran en el historial."],
        ["Almacen.importar_csv(ruta, reemplazar)", "Detecta ; o ,, valida todas las filas y carga todo o nada."],
        ["graficos.dibujar(figura, tipo, …, resaltar)", "Limpia la figura y dibuja el tipo pedido; resalta un dato."],
        ["Aplicacion.refrescar()", "Relee la base y actualiza tabla, KPI, historial y gráfico."],
    ], [6.5, 9.5])

    d.add_heading("10. Portabilidad: construcción y ejecución", 1)
    parrafo(d, "**¿Puede ejecutarse en otro equipo sin configuración compleja?** Sí. La carpeta de entrega contiene "
               "VisualizadorEncuestas.exe (unos 39 MB), LEEME.txt y un CSV de ejemplo. Se copia a una USB o a cualquier "
               "carpeta y se abre con doble clic. La primera vez crea la carpeta datos/ con encuesta.db junto al "
               "ejecutable, de modo que los datos viajan con el programa. No escribe en el registro de Windows ni "
               "instala nada.")
    parrafo(d, "**Construir el ejecutable** (solo lo necesita el desarrollador): en Windows con Python 3.12, ejecutar "
               "construir_windows.bat. Este crea un entorno virtual, instala requirements.txt, corre las pruebas y ejecuta:")
    parrafo(d, "pyinstaller --noconfirm --onefile --windowed --name VisualizadorEncuestas --paths app app\\main.py",
            cursiva=True, alinear=WD_ALIGN_PARAGRAPH.LEFT)
    parrafo(d, "Alternativamente, el workflow de GitHub Actions construye el .exe en cada push y lo publica como "
               "artefacto descargable en la pestaña Actions del repositorio.")
    parrafo(d, f"**Código fuente y documentación:** {REPO}", alinear=WD_ALIGN_PARAGRAPH.LEFT)

    d.add_heading("11. Pruebas y resultados", 1)
    parrafo(d, "**¿Cómo se comprobó la ejecución en diferentes entornos?** El código se probó en Linux (Python 3.14) y "
               "en Windows emulado con Wine (Python 3.12). El .exe final se probó en un **entorno Windows limpio, sin "
               "Python instalado**: abrió la ventana, creó datos/encuesta.db junto al ejecutable y generó los gráficos. "
               "Antes de la sustentación se recomienda abrirlo una vez en un equipo de la sala.")
    d.add_heading("11.1 Pruebas automáticas (unittest)", 2)
    tabla(d, ["Grupo", "Qué se verifica", "Resultado"], [
        ["Validación (7)", "Nombres vacíos, cortos, largos, sin letras o con símbolos; duplicados con tildes; enteros; "
                           "coherencia; población", "OK"],
        ["Indicadores (4)", "Respuestas de la FASE 1 (moda, mínimo, promedio 16,67), empates, lista vacía, Pareto "
                            "28-50-68-83-93-100 %", "OK"],
        ["Almacén (9)", "Creación con datos de la guía, persistencia al reabrir, coherencia, historial, restaurar, CSV de "
                        "ida y vuelta, CSV con errores y CSV que supera la población", "OK"],
        ["Gráficos (3)", "Los 5 tipos se generan; las alturas de las barras son los datos; mensaje sin datos", "OK"],
        ["**Total**", "**23 pruebas**", "**23 OK**"],
    ], [3.2, 10.3, 2.5])
    d.add_heading("11.2 Pruebas de la interfaz", 2)
    tabla(d, ["#", "Acción", "Resultado esperado", "Obtenido"], [
        ["1", "Agregar con nombre vacío", "No guarda; mensaje junto al campo", "OK"],
        ["2", "Agregar «Dibujar» con 5 (ya hay 100/100)", "No guarda; «Sobran 5»", "OK"],
        ["3", "Agregar «LEÉR»", "No guarda; ya existe «Leer»", "OK"],
        ["4", "Cambiar población a 90", "No guarda; menor que lo registrado", "OK"],
        ["5", "Población 120 y agregar «Dibujar o pintar» 12", "Guarda; tabla, KPI y gráfico actualizados; barra naranja", "OK"],
        ["6", "Seleccionar «Leer» y cambiar a 14", "Actualiza; KPI 116/120; queda en historial «10 → 14»", "OK"],
        ["7", "Cambiar entre los 5 tipos de gráfico", "Cada gráfico se dibuja con los datos actuales", "OK"],
        ["8", "Abrir el .exe en Windows sin Python", "Abre y crea datos/encuesta.db junto al .exe", "OK"],
    ], [0.8, 5.6, 7.3, 2.3])
    d.add_heading("11.3 Mejoras realizadas a partir de las pruebas", 2)
    vinetas(d, [
        "La barra de estado desaparecía cuando la ventana era pequeña. Se cambió el orden de empaquetado para que "
        "siempre sea visible.",
        "La nota «n = … estudiantes» se montaba sobre el título del eje X. Se movió debajo del título del gráfico.",
        "En el Pareto, los porcentajes acumulados tapaban los números de las barras. Se les puso fondo blanco y "
        "se separaron de la línea.",
        "El mensaje de duplicado repetía lo que escribió el usuario («LEÉR»). Ahora nombra la actividad existente "
        "(«Leer») y sugiere seleccionarla en la tabla.",
        "En el gráfico circular, el porcentaje de la parte «Sin registrar» (gris claro) era blanco sobre gris. Ahora "
        "usa texto oscuro.",
    ])

    d.add_heading("12. Manual de usuario breve", 1)
    vinetas(d, [
        "Abrir **VisualizadorEncuestas.exe**. Si Windows SmartScreen advierte, elegir «Más información» → «Ejecutar de "
        "todas formas» (el .exe no está firmado).",
        "Para agregar, escribir la actividad y el número de estudiantes y pulsar **Agregar** (o Enter).",
        "Para corregir, hacer clic en la fila, cambiar el dato y pulsar **Actualizar**. Para borrar, usar **Eliminar**.",
        "Elegir el **tipo de gráfico** en la lista de arriba del gráfico. **Exportar PNG** guarda la imagen.",
        "**Archivo → Importar CSV** carga datos (columnas actividad;estudiantes). **Exportar** guarda en CSV.",
        "**Archivo → Restaurar datos originales** vuelve a los datos de la guía. El historial se conserva.",
    ], "List Number")

    d.add_heading("13. Respuestas a las preguntas orientadoras", 1)
    tabla(d, ["Pregunta", "Respuesta"], [
        ["¿Qué datos se ingresan y cuál es su propósito?", "Actividad y número de estudiantes, más título y población; "
                                                           "sirven para interpretar preferencias con gráficos y KPI."],
        ["¿Cómo se garantiza que sean completos, válidos y consistentes?", "Filtro de teclado, reglas de modelo.py y "
                                                                           "restricciones de SQLite (sección 6)."],
        ["¿Por qué se seleccionó cada visualización?", "Ver la tabla de la sección 7."],
        ["¿Cómo facilita la interpretación?", "Valores escritos, orden, porcentajes sobre la población, resaltado del dato "
                                              "nuevo y KPI."],
        ["¿Se ejecuta en otro equipo sin configuración compleja?", "Sí: un .exe sin instalación; los datos quedan en "
                                                                   "datos/ junto a él."],
        ["¿Cómo se comprueba en diferentes entornos?", "Linux, Windows (Wine) con Python y Windows limpio sin Python."],
        ["¿Qué pruebas se realizaron?", "23 pruebas automáticas y 8 pruebas de interfaz (sección 11)."],
        ["¿Qué resultados y mejoras hubo?", "Todas las pruebas pasaron después de 5 mejoras de interfaz (sección 11.3)."],
    ], [6, 10])
    d.save(ENTREGA / "Fase2_Documentacion_Software.docx")


if __name__ == "__main__":
    ENTREGA.mkdir(exist_ok=True)
    fase1()
    fase2()
    print("Documentos en", ENTREGA)
